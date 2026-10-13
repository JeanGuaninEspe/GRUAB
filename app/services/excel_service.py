"""
Servicio de generación de reportes ejecutivos en formato Excel (.xlsx) con openpyxl.
Aplica estándares de diseño institucional, estilos corporativos, cabeceras del MTOP y autoajuste de columnas.
"""
import io
from datetime import datetime
from decimal import Decimal
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.schemas.reportes import ServicioReporteEnum

MESES_ES = [
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"
]


class ExcelReportService:
    def generar_reporte_mtop_excel(
        self,
        df: pd.DataFrame,
        servicio: ServicioReporteEnum,
        mes: int,
        anio: int,
    ) -> tuple[io.BytesIO, str]:
        """
        Genera un archivo Excel en memoria formateado profesionalmente para el MTOP.
        Retorna una tupla (buffer_bytes, nombre_archivo).
        """
        wb = Workbook()
        ws = wb.active
        ws.title = f"MTOP_{servicio.value.upper()}"
        ws.views.sheetView[0].showGridLines = True

        nombre_mes = MESES_ES[mes - 1] if 1 <= mes <= 12 else str(mes)
        fecha_emision = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # ----------------------------------------------------
        # DEFINICIÓN DE ESTILOS
        # ----------------------------------------------------
        font_title_inst = Font(name="Calibri", size=14, bold=True, color="1B365D")
        font_subtitle = Font(name="Calibri", size=11, bold=True, color="4A5568")
        font_report_title = Font(name="Calibri", size=12, bold=True, color="003366")
        font_meta = Font(name="Calibri", size=9, italic=True, color="718096")

        font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

        font_data = Font(name="Calibri", size=10, color="1A202C")
        fill_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        thin_border_side = Side(border_style="thin", color="D1D5DB")
        thin_border = Border(
            left=thin_border_side,
            right=thin_border_side,
            top=thin_border_side,
            bottom=thin_border_side,
        )

        header_border_side = Side(border_style="medium", color="0F243E")
        header_border = Border(
            left=thin_border_side,
            right=thin_border_side,
            top=header_border_side,
            bottom=header_border_side,
        )

        # ----------------------------------------------------
        # CABECERA INSTITUCIONAL
        # ----------------------------------------------------
        ws.merge_cells("A1:H1")
        ws["A1"] = "MINISTERIO DE TRANSPORTE Y OBRAS PÚBLICAS (MTOP)"
        ws["A1"].font = font_title_inst
        ws["A1"].alignment = align_left

        ws.merge_cells("A2:H2")
        ws["A2"] = "CONCESIÓN VIAL - SISTEMA DE GESTIÓN Y OPERACIONES DE RESCATE"
        ws["A2"].font = font_subtitle
        ws["A2"].alignment = align_left

        ws.merge_cells("A3:H3")
        ws["A3"] = f"INFORME MENSUAL DE REGISTRO DE {servicio.value.upper()}"
        ws["A3"].font = font_report_title
        ws["A3"].alignment = align_left

        ws.merge_cells("A4:H4")
        ws["A4"] = f"Periodo: {nombre_mes} {anio}  |  Fecha de Emisión: {fecha_emision}  |  Total Registros: {len(df)}"
        ws["A4"].font = font_meta
        ws["A4"].alignment = align_left

        # Fila en blanco
        start_row = 6

        # ----------------------------------------------------
        # ENCABEZADOS DE LA TABLA
        # ----------------------------------------------------
        columns = list(df.columns) if not df.empty else ["SIN DATOS"]

        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=start_row, column=col_idx)
            header_text = str(col_name).replace("_", " ").upper()
            cell.value = header_text
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = header_border

        ws.row_dimensions[start_row].height = 28

        # ----------------------------------------------------
        # CUERPO DE DATOS
        # ----------------------------------------------------
        current_row = start_row + 1

        if df.empty:
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max(len(columns), 4))
            empty_cell = ws.cell(row=current_row, column=1)
            empty_cell.value = "No se encontraron registros para el periodo seleccionado."
            empty_cell.font = Font(name="Calibri", size=10, italic=True, color="718096")
            empty_cell.alignment = align_center
            ws.row_dimensions[current_row].height = 24
        else:
            for row_idx, record in enumerate(df.to_dict(orient="records"), start=0):
                row_fill = fill_even if row_idx % 2 == 0 else fill_odd
                ws.row_dimensions[current_row].height = 20

                for col_idx, col_name in enumerate(columns, start=1):
                    val = record.get(col_name)
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.fill = row_fill
                    cell.font = font_data
                    cell.border = thin_border

                    # Formateo de tipos de datos
                    if isinstance(val, (int,)):
                        cell.value = val
                        cell.alignment = align_right
                        cell.number_format = "#,##0"
                    elif isinstance(val, (float, Decimal)):
                        cell.value = float(val)
                        cell.alignment = align_right
                        cell.number_format = "#,##0.00"
                    elif isinstance(val, (datetime, pd.Timestamp)):
                        cell.value = val.strftime("%Y-%m-%d %H:%M:%S")
                        cell.alignment = align_center
                    elif isinstance(val, pd.Series) or val is None or pd.isna(val):
                        cell.value = ""
                        cell.alignment = align_center
                    elif isinstance(val, bool):
                        cell.value = "SI" if val else "NO"
                        cell.alignment = align_center
                    else:
                        cell.value = str(val)
                        cell.alignment = align_left

                current_row += 1

        # ----------------------------------------------------
        # AUTOAJUSTE DE ANCHOS DE COLUMNAS
        # ----------------------------------------------------
        for col_idx, col_name in enumerate(columns, start=1):
            col_letter = get_column_letter(col_idx)
            header_len = len(str(col_name).replace("_", " "))
            max_len = header_len

            # Revisar las primeras 100 filas para calcular ancho
            sample_rows = df[col_name].astype(str).head(100) if not df.empty and col_name in df else []
            for item in sample_rows:
                if item and item != "None" and item != "nan":
                    max_len = max(max_len, len(item))

            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        # ----------------------------------------------------
        # GUARDAR EN BUFFER BINARIO
        # ----------------------------------------------------
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"reporte_mtop_{servicio.value}_{anio}_{mes:02d}.xlsx"
        return output, filename


excel_service = ExcelReportService()
