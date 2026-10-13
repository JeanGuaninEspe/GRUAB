"""
Pruebas unitarias para el generador de reportes Excel openpyxl.
"""
from decimal import Decimal
import pandas as pd
from openpyxl import load_workbook
from app.services.excel_service import excel_service
from app.schemas.reportes import ServicioReporteEnum


def test_generar_excel_con_datos():
    data = [
        {
            "codigo_formulario": "AMB-2026-001",
            "fecha": "2026-08-01",
            "turno": 1,
            "abscisa_km": Decimal("25.500"),
            "sentido": "ASCENDENTE",
            "nombre_paciente": "Juan Perez",
            "motivo_atencion": "Colisión",
            "severidad": "LEVE",
            "km_recorridos": Decimal("45.50"),
        },
        {
            "codigo_formulario": "AMB-2026-002",
            "fecha": "2026-08-02",
            "turno": 2,
            "abscisa_km": Decimal("12.300"),
            "sentido": "DESCENDENTE",
            "nombre_paciente": "Maria Gomez",
            "motivo_atencion": "Volcamiento",
            "severidad": "GRAVE",
            "km_recorridos": Decimal("30.00"),
        },
    ]
    df = pd.DataFrame(data)

    stream, filename = excel_service.generar_reporte_mtop_excel(
        df=df,
        servicio=ServicioReporteEnum.AMBULANCIA,
        mes=8,
        anio=2026,
    )

    assert filename == "reporte_mtop_ambulancia_2026_08.xlsx"
    assert stream.getbuffer().nbytes > 0

    # Cargar y verificar estructura del archivo generado
    wb = load_workbook(stream)
    ws = wb.active
    assert ws.title == "MTOP_AMBULANCIA"
    assert "MINISTERIO DE TRANSPORTE Y OBRAS PÚBLICAS" in ws["A1"].value
    assert ws.max_row >= 8  # 4 filas de banner + 1 blanca + 1 cabecera + 2 datos


def test_generar_excel_sin_datos():
    df = pd.DataFrame()
    stream, filename = excel_service.generar_reporte_mtop_excel(
        df=df,
        servicio=ServicioReporteEnum.GRUA,
        mes=5,
        anio=2026,
    )

    assert filename == "reporte_mtop_grua_2026_05.xlsx"
    wb = load_workbook(stream)
    ws = wb.active
    assert ws.title == "MTOP_GRUA"
